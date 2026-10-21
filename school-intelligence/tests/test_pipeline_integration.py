"""Integration, end-to-end, regression and negative tests over a local HTTP fixture site.

Covers: adapter -> extraction -> verification -> JSON -> Excel -> independent validation.
"""

import json

import pytest
from openpyxl import load_workbook

from masteredge import validate_final
from masteredge.pipeline import run
from tests.conftest import greenfield_seed, make_project


def master(root):
    return json.loads((root / "data" / "output" / "schools_master.json").read_text())


def fact(rec, name):
    return rec["fields"].get(name) or {}


@pytest.fixture()
def slice1(site, tmp_path):
    base, _ = site
    root = make_project(tmp_path, [greenfield_seed(base)])
    run(1, root=root)
    return root, base


def test_e2e_official_site_extraction(slice1):
    root, base = slice1
    rec = master(root)["schools"][0]
    fee = fact(rec, "annual_fee")
    assert fee["status"] == "official_source"
    assert (fee["value"]["min"], fee["value"]["max"], fee["value"]["academic_year"], fee["value"]["period"]) == (450000, 720000, "2025-26", "annual")
    assert fact(rec, "admission_charges")["value"]["amount"] == 5000
    codes = {i["value"]["code"] for i in fact(rec, "curricula")["items"]}
    assert {"CBSE", "CAMBRIDGE_IGCSE", "IB_DP"} <= codes
    assert "CAMBRIDGE_AS_A_LEVEL" not in codes and "ISC" not in codes        # blog page ignored
    assert "CISCE" not in codes and "ICSE" not in codes                      # script text ignored
    assert fact(rec, "founding_year")["value"] == 1998                       # blog 'Founded in 1950' ignored
    assert fact(rec, "campus_student_count")["value"] == 1800                # webinar 9000 ignored
    assert fact(rec, "average_class_size")["value"] == 24
    assert fact(rec, "student_teacher_ratio")["value"] == "12:1"
    assert {"board": "CBSE", "type": "affiliation_number", "id": "1130123"} in fact(rec, "affiliation_details")["value"]
    contacts = {c["value"] for c in fact(rec, "public_contact")["value"]}
    assert "9876543210" not in contacts and "ramesh.k@greenfieldschool.in" not in contacts
    assert fact(rec, "grades_offered")["value"] == {"grade_min": -2, "grade_max": 12}
    # robots-disallowed page never fetched: its fee concession must not appear anywhere
    assert not any("/private/" in e["source_url"] for e in rec["evidence"])
    # every evidence excerpt really occurs in its stored source
    assert all(e["verification_status"] != "unverified" for e in rec["evidence"])
    # broken PDF was attempted, logged and did not abort the crawl
    assert rec["research_checks"]["official_site_crawl"]["status"] == "done"
    assert rec["analysis"]["premium_category"]["value"] == "premium"
    assert not rec["analysis"]["premium_category"]["provisional"]


def test_e2e_metrics_and_excel(slice1):
    root, _ = slice1
    vr = json.loads((root / "data" / "output" / "validation_report.json").read_text())
    m = vr["metrics"]
    assert vr["schema_valid"] is True
    assert m["fee_verification_pct"] == 100.0 and m["enrollment_verification_pct"] == 100.0
    assert m["curriculum_verification_pct"] == 0.0           # no registry confirmation possible here
    wb = load_workbook(root / "data" / "output" / "MASTEREDGE_India_Premium_School_Opportunity_Report.xlsx")
    for sheet in validate_final.REQUIRED_SHEETS + ["Evidence", "Source Registry"]:
        assert sheet in wb.sheetnames
    ws = wb["Top 100 Schools"]
    hdr = [c.value for c in ws[1]]
    row = [c for c in ws[2]]
    assert row[hdr.index("Max annual fee (INR)")].value == 720000
    assert row[hdr.index("Fee source")].hyperlink is not None


def test_independent_validation_requires_manual_audit(slice1):
    root, _ = slice1
    rep = validate_final.validate(root=root)
    assert rep["release_status"] == "FAILED_VALIDATION"
    failed = {r["id"] for r in rep["rules"] if r["status"] == "FAIL"}
    assert failed == {"R16_audit"}
    # record MATCH verdicts for the sample -> passes (registry target remains unmet)
    sample = rep["manual_audit"]["sample"]
    (root / "data" / "audit").mkdir(parents=True, exist_ok=True)
    (root / "data" / "audit" / "manual_audit.json").write_text(json.dumps({"results": [
        {"audit_key": s["audit_key"], "verdict": "MATCH", "note": "fixture"} for s in sample]}))
    rep2 = validate_final.validate(root=root)
    assert rep2["release_status"] == "PASSED_WITH_UNMET_TARGETS"
    assert rep2["automated_recheck"]["automated_accuracy_pct"] == 100.0


def test_validator_is_read_only(slice1):
    root, _ = slice1
    out = root / "data" / "output"
    before = {p.name: p.read_bytes() for p in out.iterdir() if p.is_file()}
    rep = validate_final.validate(root=root)
    after = {p.name: p.read_bytes() for p in out.iterdir() if p.is_file()}
    assert before == after
    assert next(r for r in rep["rules"] if r["id"] == "R00_read_only")["status"] == "PASS"


def test_validator_detects_tampered_value(slice1):
    root, _ = slice1
    p = root / "data" / "output" / "schools_master.json"
    m = json.loads(p.read_text())
    m["schools"][0]["fields"]["annual_fee"]["value"]["max"] = 999999
    for e in m["schools"][0]["evidence"]:
        if e["field_name"] == "annual_fee":
            e["normalized_value"]["max"] = 999999
    p.write_text(json.dumps(m))
    rep = validate_final.validate(root=root)
    assert rep["release_status"] == "FAILED_VALIDATION"
    assert any(f["rule"] in ("R04_fees", "R10_report_numbers", "R08_premium_rules") for f in rep["failures"])


# ------------------------------------------------------------------ regression

def test_rerun_is_stable_and_preserves_evidence(slice1):
    root, base = slice1
    m1 = master(root)
    ev1 = {e["evidence_id"] for e in m1["schools"][0]["evidence"]}
    v1 = m1["schools"][0]["record_version"]
    run(1, root=root, refresh=True)
    m2 = master(root)
    assert [s["school_id"] for s in m2["schools"]] == [s["school_id"] for s in m1["schools"]]  # no duplicate campus
    ev2 = {e["evidence_id"] for e in m2["schools"][0]["evidence"]}
    assert ev1 <= ev2                                          # historical evidence preserved
    assert m2["schools"][0]["record_version"] == v1            # unchanged content -> same version
    assert m2["schools"][0]["fields"]["annual_fee"] == m1["schools"][0]["fields"]["annual_fee"]


def test_verified_fee_not_silently_removed_when_page_changes(site, tmp_path):
    base, site_dir = site
    root = make_project(tmp_path, [greenfield_seed(base)])
    run(1, root=root)
    (site_dir / "fees" / "index.html").write_text("<html><head><title>Fees</title></head><body><p>Contact us for fees.</p></body></html>")
    rc = json.loads((root / "config" / "research_config.json").read_text())
    rc["refresh_interval_days"] = 0          # force refetch
    (root / "config" / "research_config.json").write_text(json.dumps(rc))
    run(1, root=root, refresh=True)
    rec = master(root)["schools"][0]
    fee = rec["fields"]["annual_fee"]
    assert fee["status"] == "needs_review" and fee["value"]["max"] == 720000
    assert any(e["field_name"] == "annual_fee" and e["superseded"] for e in rec["evidence"])  # kept, flagged


# ------------------------------------------------------------------ third-party + negatives

def test_third_party_conflicts_are_flagged_not_resolved(site, tmp_path):
    base, _ = site
    root = make_project(tmp_path, [greenfield_seed(base, with_profile=True)])
    run(2, root=root)
    rec = master(root)["schools"][0]
    fy = rec["fields"]["founding_year"]
    assert fy["status"] == "conflict" and fy["value"] == 1998, (rec["research_checks"].get("third_party_crosscheck"), rec["research_checks"].get("third_party_fee"), len(rec["evidence"]), [e["source_url"] for e in rec["evidence"] if e["field_name"] == "founding_year"])
    assert {v["value"] for v in fy["conflict"]["values"]} == {1998, 1999}
    fee = rec["fields"]["annual_fee"]
    assert fee["status"] == "conflict" and fee["value"]["max"] == 720000
    assert any(g["field"] == "founding_year" and g["gap_type"] == "conflict" for g in rec["data_gaps"])
    # registry sources are disabled in tests -> registry check is recorded as blocked, never as done
    assert rec["research_checks"]["registry_curriculum_check"]["status"] == "blocked"
    wb = load_workbook(root / "data" / "output" / "MASTEREDGE_India_Premium_School_Opportunity_Report.xlsx")
    gaps = {(r[0].value, r[2].value, r[3].value) for r in wb["Data Gaps"].iter_rows(min_row=2)}
    assert (rec["school_id"], "founding_year", "conflict") in gaps
    rep = validate_final.validate(root=root)
    assert next(r for r in rep["rules"] if r["id"] == "R13_conflicts")["status"] == "PASS"


def test_unavailable_website_is_reported_not_fabricated(tmp_path):
    seed = {"name": "Nowhere Academy", "campus": "Central", "city": "Pune", "state": "Maharashtra",
            "official_website": "http://127.0.0.1:9/"}
    root = make_project(tmp_path, [seed])
    run(1, root=root)
    rec = master(root)["schools"][0]
    assert rec["research_checks"]["official_website_identity"]["status"] in ("failed", "blocked")
    assert rec["research_checks"]["fee_research"]["status"] == "failed"
    assert all(f.get("value") in (None, []) for k, f in rec["fields"].items() if k != "premium_category")
    assert rec["analysis"]["premium_category"]["value"] == "unclassified"


def test_wrong_website_identity_rejected(site, tmp_path):
    base, _ = site
    seed = {"name": "Riverside Montessori", "campus": "Kothrud", "city": "Pune", "state": "Maharashtra", "official_website": base}
    root = make_project(tmp_path, [seed])
    run(1, root=root)
    rec = master(root)["schools"][0]
    assert rec["research_checks"]["official_website_identity"]["status"] == "failed"
    assert rec["fields"]["annual_fee"]["value"] is None   # facts from someone else's site never attached


def test_missing_curriculum_history_reported(site, tmp_path):
    """History needs the Wayback service; when it is unreachable the check fails visibly."""
    base, _ = site
    root = make_project(tmp_path, [greenfield_seed(base)], slices_override={"4": {"discovery": False, "third_party": False, "registry_checks": False}})
    src = json.loads((root / "config" / "sources.json").read_text())
    (root / "config" / "sources.json").write_text(json.dumps(src))
    import masteredge.adapters.wayback as wb_mod
    orig = wb_mod.CDX
    wb_mod.CDX = "http://127.0.0.1:9/cdx"           # unreachable archive
    try:
        run(4, root=root, limit=1)
    finally:
        wb_mod.CDX = orig
    rec = master(root)["schools"][0]
    assert rec["research_checks"]["curriculum_history_research"]["status"] in ("failed", "not_publicly_available")
    assert rec["fields"]["curriculum_history"]["value"] == []


def test_cisce_registry_lookup_with_campus_disambiguation(site, tmp_path):
    from masteredge.adapters import registries
    from masteredge.config import load_config
    from masteredge.pipeline import build_context
    from masteredge.records import new_school_record

    base, _ = site
    root = make_project(tmp_path, [])
    ctx = build_context(load_config(root))
    ctx.source_status = {"cisce": {"access": "ACCESSIBLE"}}
    registries._cisce_cache.clear()
    orig = registries.CISCE_RESULT
    registries.CISCE_RESULT = base + "cisce/result/"
    try:
        rec = new_school_record("IN-riverside-model-school-kothrud-pune-000001",
                                {"name": "Riverside Model School", "campus": "Kothrud", "city": "Pune", "state": "Maharashtra"}, {})
        r = registries.cisce_lookup(ctx, rec, max_pages=5)
        assert r["check"]["status"] == "done" and r["check"]["matched"]
        got = {(e["field_name"], json.dumps(e["normalized_value"], sort_keys=True)) for e in r["evidence"]}
        assert ("affiliation_details", json.dumps({"board": "CISCE", "id": "MH999", "type": "school_code"}, sort_keys=True)) in got
        assert all(e["verification_status"] == "official_registry_verified" for e in r["evidence"])
        # same name, no campus -> ambiguous -> nothing asserted
        registries._cisce_cache.clear()
        rec2 = new_school_record("IN-riverside-model-school-pune-000002",
                                 {"name": "Riverside Model School", "city": "Pune", "state": "Maharashtra"}, {})
        r2 = registries.cisce_lookup(ctx, rec2, max_pages=5)
        assert r2["evidence"] == [] and "ambiguous" in r2["check"]["detail"]
    finally:
        registries.CISCE_RESULT = orig
        registries._cisce_cache.clear()


def test_excel_survives_control_characters(site, tmp_path):
    """Regression: scraped text with control characters crashed openpyxl in the 100-school run."""
    base, site_dir = site
    p = site_dir / "about" / "index.html"
    p.write_text(p.read_text(encoding="utf-8").replace("Our Principal", "Our\x0b Principal"), encoding="utf-8")
    root = make_project(tmp_path, [greenfield_seed(base)])
    run(1, root=root)
    assert (root / "data" / "output" / "MASTEREDGE_India_Premium_School_Opportunity_Report.xlsx").exists()
