"""CEO-ready Excel report generation (openpyxl)."""

from __future__ import annotations

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.utils import get_column_letter

from .analysis import ANALYSIS_LABEL
from .config import Config
from .normalize import grade_label
from .records import VERIFIED_FIELD_STATUSES, display_value
from .util import utc_now

HDR_FILL = PatternFill("solid", fgColor="1F3864")
HDR_FONT = Font(bold=True, color="FFFFFF")
TITLE_FONT = Font(bold=True, size=16, color="1F3864")
SUB_FONT = Font(bold=True, size=12, color="1F3864")
WARN_FILL = PatternFill("solid", fgColor="FCE4D6")
OK_FILL = PatternFill("solid", fgColor="E2EFDA")
LINK_FONT = Font(color="0563C1", underline="single")

STATUS_TEXT = {
    "official_registry_verified": "Verified (official registry)",
    "official_source": "Official school website",
    "archive_only": "Archive only",
    "third_party_only": "Third-party only (unverified)",
    "conflict": "CONFLICT - see Data Gaps",
    "not_found": "Not publicly available",
    "needs_review": "Needs review",
    "analysis": "Analysis",
}


def _clean(v):
    """Excel forbids some control characters; strip them from display text (JSON keeps the raw text)."""
    if isinstance(v, str):
        return ILLEGAL_CHARACTERS_RE.sub("", v)
    return v


def _append(ws, row):
    ws.append([_clean(v) for v in row])


def _sheet(wb, title, headers, widths=None):
    ws = wb.create_sheet(title)
    _append(ws, headers)
    for i, _ in enumerate(headers, 1):
        c = ws.cell(row=1, column=i)
        c.fill, c.font = HDR_FILL, HDR_FONT
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.column_dimensions[get_column_letter(i)].width = (widths or {}).get(i, 18)
    ws.freeze_panes = "B2"
    return ws


def _link(ws, row, col, url, text=None):
    c = ws.cell(row=row, column=col, value=_clean(text or url))
    if url and str(url).startswith("http"):
        c.hyperlink = url
        c.font = LINK_FONT
    return c


def _wrap(ws):
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")


def _first_source(rec, fact) -> str | None:
    if not fact:
        return None
    ev = {e["evidence_id"]: e for e in rec["evidence"]}
    for i in fact.get("evidence_ids", []):
        if i in ev:
            return ev[i]["source_url"]
    return None


def _status(fact) -> str:
    if not fact:
        return STATUS_TEXT["not_found"]
    return STATUS_TEXT.get(fact.get("status"), fact.get("status"))


def ranked(schools):
    return sorted(schools, key=lambda r: (-r["analysis"]["masteredge_fit"]["score"], r["school_id"]))


def school_label(rec) -> tuple[str, str]:
    f = rec["fields"]
    name = (f.get("school_name") or {}).get("value") or rec["input"]["name"]
    campus = (f.get("campus") or {}).get("value") or rec["input"].get("campus") or ""
    return name, campus


def build_report(cfg: Config, schools: list[dict], groups: list[dict], metrics: dict, sources: list[dict]) -> str:
    wb = Workbook()
    wb.remove(wb.active)
    rk = ranked(schools)
    gates = {g["id"]: g for g in cfg.validation_rules["gates"]}

    # 1. CEO Dashboard ----------------------------------------------------------
    ws = wb.create_sheet("CEO Dashboard")
    ws.column_dimensions["A"].width = 48
    for col in "BCDEF":
        ws.column_dimensions[col].width = 22
    ws["A1"] = "MASTEREDGE — India Premium School Opportunity Report"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = f"Generated {utc_now()} · {len(schools)} schools researched · every fact links to evidence (see Evidence sheet)"
    cats = {}
    for r in schools:
        pc = r["analysis"]["premium_category"]
        cats.setdefault(pc["value"], []).append(r)
    row = 4
    ws.cell(row=row, column=1, value="Headline numbers").font = SUB_FONT
    row += 1
    kpis = [
        ("Schools researched", len(schools)),
        ("Ultra-premium by published fee", len(cats.get("ultra_premium", []))),
        ("Premium by published fee", len(cats.get("premium", []))),
        ("…of which provisional (third-party fee or fee period not stated)",
         sum(1 for r in schools if r["analysis"]["premium_category"]["value"] in ("ultra_premium", "premium")
             and r["analysis"]["premium_category"]["provisional"])),
        ("…of which category conflicts between sources",
         sum(1 for r in schools if r["analysis"]["premium_category"].get("conflict"))),
        ("Unclassified (no published fee found)", len(cats.get("unclassified", []))),
        ("Premium school groups (shared official domain)", len(groups)),
        ("Fields with unresolved source conflicts", metrics["fields_in_conflict"]),
    ]
    for k, v in kpis:
        ws.cell(row=row, column=1, value=k)
        ws.cell(row=row, column=2, value=v)
        row += 1
    row += 1
    ws.cell(row=row, column=1, value="Data quality (actual, recalculated from final JSON)").font = SUB_FONT
    row += 1
    for i, h in enumerate(["Metric", "Actual %", "Target", "Status"], 1):
        c = ws.cell(row=row, column=i, value=h)
        c.fill, c.font = HDR_FILL, HDR_FONT
    row += 1
    for key, label in [("field_coverage_pct", "Required-field coverage"),
                       ("evidence_coverage_pct", "Evidence coverage (populated facts with valid source)"),
                       ("official_source_verification_pct", "Official-source verification (facts)"),
                       ("curriculum_verification_pct", "Curriculum verified by official registry"),
                       ("curriculum_official_site_pct", "Curriculum stated on official website"),
                       ("fee_verification_pct", "Fee verified & dated (official)"),
                       ("fee_official_any_pct", "Any fee published on official website"),
                       ("enrollment_verification_pct", "Enrollment verified with reporting year"),
                       ("research_completion_pct", "Research completion")]:
        gate_id = {"field_coverage_pct": "required_field_coverage_pct"}.get(key, key)
        g = gates.get(gate_id, {})
        val = metrics.get(key)
        tgt = g.get("target")
        st = "report only" if tgt is None else ("MET" if val is not None and val >= tgt else "NOT MET")
        ws.cell(row=row, column=1, value=label)
        ws.cell(row=row, column=2, value=val)
        ws.cell(row=row, column=3, value=tgt)
        c = ws.cell(row=row, column=4, value=st)
        c.fill = OK_FILL if st == "MET" else (WARN_FILL if st == "NOT MET" else PatternFill())
        row += 1
    row += 1
    ws.cell(row=row, column=1, value="Top prospects (MASTEREDGE fit = ANALYSIS, not a purchase prediction)").font = SUB_FONT
    row += 1
    for i, h in enumerate(["School", "City", "Premium category", "Fit score /100", "Evidence completeness %", "Official website"], 1):
        c = ws.cell(row=row, column=i, value=h)
        c.fill, c.font = HDR_FILL, HDR_FONT
    row += 1
    for r in rk[:10]:
        name, campus = school_label(r)
        ws.cell(row=row, column=1, value=f"{name}{' — ' + campus if campus else ''}")
        ws.cell(row=row, column=2, value=r["input"].get("city"))
        pc = r["analysis"]["premium_category"]
        ws.cell(row=row, column=3, value=pc["label"] + (" (provisional)" if pc["provisional"] else "") + (" [CONFLICT]" if pc.get("conflict") else ""))
        ws.cell(row=row, column=4, value=r["analysis"]["masteredge_fit"]["score"])
        ws.cell(row=row, column=5, value=r["analysis"]["masteredge_fit"]["evidence_completeness_pct"])
        _link(ws, row, 6, r.get("official_website_url"))
        row += 1
    row += 1
    ws.cell(row=row, column=1, value="Source access (live check this run)").font = SUB_FONT
    row += 1
    for s in sources:
        ws.cell(row=row, column=1, value=s["name"])
        c = ws.cell(row=row, column=2, value=s.get("access"))
        c.fill = OK_FILL if s.get("access") == "ACCESSIBLE" else WARN_FILL
        ws.cell(row=row, column=3, value=s.get("reason"))
        row += 1
    row += 1
    ws.cell(row=row, column=1, value="How to read this report").font = SUB_FONT
    for line in [
        "Premium category uses the maximum published annual fee from the highest-tier source (config/premium_rules.json).",
        "‘Official school website’ = stated by the school itself; ‘Verified (official registry)’ = confirmed on an IB/CBSE/CISCE/Cambridge registry page.",
        "CONFLICT means sources disagree; both values are kept and listed in Data Gaps. Nothing is silently resolved.",
        "Missing numbers are blank, never estimated. Fit scores are analysis built only from evidenced facts.",
    ]:
        row += 1
        ws.cell(row=row, column=1, value=line)

    # 2. Top 100 Schools -----------------------------------------------------------
    ws = _sheet(wb, "Top 100 Schools", ["Rank", "School ID", "School", "Campus", "City", "State", "Premium category",
                                         "Category flags", "Annual fee (displayed)", "Fee status", "Max annual fee (INR)",
                                         "Curricula", "Curriculum status", "Grades", "Campus students", "Enrollment status",
                                         "Fit score /100 (ANALYSIS)", "Evidence completeness %", "Why (evidence-based)",
                                         "Fee source", "Curriculum source"],
                {3: 34, 9: 30, 12: 26, 19: 60, 20: 40, 21: 40})
    for i, r in enumerate(rk[: cfg.research["report"]["top_n"]], 1):
        f = r["fields"]
        name, campus = school_label(r)
        pc = r["analysis"]["premium_category"]
        flags = []
        if pc["provisional"]:
            flags.append("provisional: " + "; ".join(pc.get("provisional_reasons", [])))
        if pc.get("conflict"):
            flags.append("CONFLICT: " + pc["conflict"]["summary"])
        fee = f.get("annual_fee") or {}
        cnt = f.get("campus_student_count") or {}
        comps = r["analysis"]["masteredge_fit"]["components"]
        why = "; ".join(f"{k.replace('_', ' ')} {v['score']:.2f}: {v['explanation']}" for k, v in comps.items() if v["score"] > 0)
        _append(ws, [i, r["school_id"], name, campus, r["input"].get("city"), r["input"].get("state"), pc["label"],
                   " | ".join(flags), display_value("annual_fee", fee), _status(fee),
                   (fee.get("value") or {}).get("max"), display_value("curricula", f.get("curricula")),
                   _status(f.get("curricula")), display_value("grades_offered", f.get("grades_offered")),
                   cnt.get("value"), _status(cnt), r["analysis"]["masteredge_fit"]["score"],
                   r["analysis"]["masteredge_fit"]["evidence_completeness_pct"], why or "no positive evidence"])
        rown = ws.max_row
        _link(ws, rown, 20, _first_source(r, fee))
        _link(ws, rown, 21, _first_source(r, f.get("curricula")))
        if fee.get("status") == "conflict" or pc.get("conflict"):
            ws.cell(row=rown, column=9).fill = WARN_FILL
    _wrap(ws)

    # 3. Premium School Groups -------------------------------------------------------
    ws = _sheet(wb, "Premium School Groups", ["Group ID", "Brand domain", "Grouping basis", "Campuses in dataset",
                                               "Member schools", "Member premium categories", "Parent organisation (stated)",
                                               "Stated number of campuses", "Group student count (stated)", "Note"],
                {5: 60, 6: 40, 7: 40, 10: 50})
    by_id = {r["school_id"]: r for r in schools}
    for g in groups:
        members = [by_id[m] for m in g["member_school_ids"] if m in by_id]
        _append(ws, [g["group_id"], g["brand_domain"], g["basis"], g["campuses_in_dataset"],
                   "\n".join(" — ".join(x for x in school_label(m) if x) for m in members),
                   "\n".join(m["analysis"]["premium_category"]["label"] for m in members),
                   "; ".join(p["name"] for p in g["parent_organization_candidates"]) or "not stated",
                   "; ".join(str(s["value"]) for s in g["stated_number_of_campuses"]) or "not stated",
                   "; ".join(f"{s['value']} ({s['status']})" for s in g["group_student_count_candidates"]) or "not stated",
                   g["note"]])
    _wrap(ws)

    # 4. School Intelligence -----------------------------------------------------------
    intel_fields = ["school_name", "campus", "city", "state", "address", "official_website", "founding_year", "school_type",
                    "grades_offered", "annual_fee", "admission_charges", "campus_student_count", "group_student_count",
                    "student_teacher_ratio", "average_class_size", "curricula", "curriculum_by_grade", "subjects_offered",
                    "teaching_methodology", "assessment_methodology", "affiliation_details"]
    headers = ["School ID", "School"]
    for fn in intel_fields:
        headers += [fn.replace("_", " ").title(), f"{fn} status", f"{fn} source"]
    ws = _sheet(wb, "School Intelligence", headers, {2: 34})
    for r in rk:
        name, campus = school_label(r)
        vals = [r["school_id"], f"{name}{' — ' + campus if campus else ''}"]
        for fn in intel_fields:
            fact = r["fields"].get(fn)
            vals += [display_value(fn, fact), _status(fact), None]
        _append(ws, vals)
        rown = ws.max_row
        for j, fn in enumerate(intel_fields):
            _link(ws, rown, 3 + j * 3 + 2, _first_source(r, r["fields"].get(fn)))
    _wrap(ws)

    # 5. Curriculum History ----------------------------------------------------------------
    ws = _sheet(wb, "Curriculum History", ["School ID", "School", "Event type", "Curriculum", "Date", "Evidence date",
                                            "Source", "Excerpt", "Caveat", "Evidence ID"], {2: 30, 7: 50, 8: 60, 9: 40})
    for r in rk:
        name, campus = school_label(r)
        ev = {e["evidence_id"]: e for e in r["evidence"]}
        hist = r["fields"].get("curriculum_history") or {}
        for item in hist.get("items") or []:
            v = item["value"]
            e = ev[item["evidence_ids"][0]]
            _append(ws, [r["school_id"], f"{name}{' — ' + campus if campus else ''}", v.get("event_type"), v.get("code"),
                       v.get("date") or v.get("year"), e.get("publication_date"), None, e["evidence_excerpt"][:500],
                       e["qualifiers"].get("caveat", ""), e["evidence_id"]])
            _link(ws, ws.max_row, 7, e["source_url"])
    _wrap(ws)

    # 6. MASTEREDGE Fit -------------------------------------------------------------------------
    comp_keys = list(cfg.research["scoring_weights"].keys())
    hdr = ["School ID", "School", "Fit score /100", "Evidence completeness %"]
    for k in comp_keys:
        hdr += [f"{k} (w={cfg.research['scoring_weights'][k]})", f"{k} explanation", f"{k} missing evidence"]
    hdr += ["Recommended grades", "Implementation complexity (ANALYSIS)", "Complexity basis", "Label"]
    ws = _sheet(wb, "MASTEREDGE Fit", hdr, {2: 34})
    for r in rk:
        fit = r["analysis"]["masteredge_fit"]
        rec = r["analysis"]["recommendations"]
        name, campus = school_label(r)
        vals = [r["school_id"], f"{name}{' — ' + campus if campus else ''}", fit["score"], fit["evidence_completeness_pct"]]
        for k in comp_keys:
            c = fit["components"][k]
            vals += [round(c["score"], 3), c["explanation"], ", ".join(c["missing_evidence"])]
        vals += [rec["recommended_grades"], rec["implementation_complexity"], rec["implementation_complexity_basis"], ANALYSIS_LABEL]
        _append(ws, vals)
    _wrap(ws)

    # 7. Top 20 Sales Targets -------------------------------------------------------------------
    ws = _sheet(wb, "Top 20 Sales Targets", ["Rank", "School ID", "School", "City", "Fit score /100 (ANALYSIS)",
                                              "Premium category", "Decision-maker roles evidenced on official site",
                                              "Suggested roles to approach (generic)", "Recommended grades",
                                              "Recommended subjects", "Implementation complexity", "Pitch angle (ANALYSIS)",
                                              "Public business contact", "Contact source", "Official website"],
                {3: 34, 7: 40, 8: 40, 12: 60, 13: 40, 14: 40})
    for i, r in enumerate(rk[: cfg.research["report"]["top_sales_targets"]], 1):
        rec = r["analysis"]["recommendations"]
        name, campus = school_label(r)
        pc = r["analysis"]["premium_category"]
        contact = r["fields"].get("public_contact")
        _append(ws, [i, r["school_id"], f"{name}{' — ' + campus if campus else ''}", r["input"].get("city"),
                   r["analysis"]["masteredge_fit"]["score"], pc["label"] + (" (provisional)" if pc["provisional"] else ""),
                   ", ".join(rec["decision_maker_roles_evidenced"]) or "none found on official pages examined",
                   ", ".join(rec["decision_maker_roles_suggested"]), rec["recommended_grades"],
                   ", ".join(rec["recommended_subjects"]), rec["implementation_complexity"], rec["pitch_angle"],
                   display_value("public_contact", contact) or "not found", None, None])
        _link(ws, ws.max_row, 14, _first_source(r, contact))
        _link(ws, ws.max_row, 15, r.get("official_website_url"))
    _wrap(ws)

    # 8. Data Gaps ---------------------------------------------------------------------------------
    ws = _sheet(wb, "Data Gaps", ["School ID", "School", "Field / check", "Gap type", "Detail", "Conflicting values"],
                {2: 34, 5: 60, 6: 70})
    for r in rk:
        name, campus = school_label(r)
        for g in r["data_gaps"]:
            vals = g.get("values")
            _append(ws, [r["school_id"], f"{name}{' — ' + campus if campus else ''}", g["field"], g["gap_type"],
                       g.get("detail"), _fmt_values(vals)])
            if g["gap_type"] == "conflict":
                ws.cell(row=ws.max_row, column=4).fill = WARN_FILL
    _wrap(ws)

    # 9. Source References -------------------------------------------------------------------------------
    ws = _sheet(wb, "Source References", ["School ID", "School", "Field", "Source URL", "Source title", "Source type",
                                           "Publication date", "Retrieved at", "Evidence count"], {2: 30, 4: 60, 5: 40})
    for r in rk:
        name, campus = school_label(r)
        seen = {}
        for e in r["evidence"]:
            if e.get("superseded"):
                continue
            k = (e["field_name"], e["source_url"])
            if k in seen:
                seen[k]["n"] += 1
                continue
            seen[k] = {"e": e, "n": 1}
        for (fname, url), d in sorted(seen.items()):
            e = d["e"]
            _append(ws, [r["school_id"], f"{name}{' — ' + campus if campus else ''}", fname, None, e["source_title"],
                       e["source_type"], e.get("publication_date"), e["retrieved_at"], d["n"]])
            _link(ws, ws.max_row, 4, url)
    _wrap(ws)

    # 10. Validation Report --------------------------------------------------------------------------------
    ws = _sheet(wb, "Validation Report", ["Metric", "Actual", "Target", "Gate status", "Definition"], {1: 44, 5: 90})
    for key, val in metrics.items():
        if key == "definitions":
            continue
        gate_id = {"field_coverage_pct": "required_field_coverage_pct"}.get(key, key)
        g = gates.get(gate_id, {})
        tgt = g.get("target")
        st = "" if tgt is None else ("MET" if val is not None and val >= tgt else "NOT MET")
        _append(ws, [key, val, tgt, st, metrics["definitions"].get(key, "")])
        if st == "NOT MET":
            ws.cell(row=ws.max_row, column=4).fill = WARN_FILL
    _append(ws, [])
    _append(ws, ["Independent validation", "See data/output/validation/final_validation_report.json and final_validation_summary.md",
               None, None, "The read-only validator recomputes every percentage above from the final JSON and checks this workbook against it."])
    _wrap(ws)

    # 11. Evidence ---------------------------------------------------------------------------------------------
    ws = _sheet(wb, "Evidence", ["Evidence ID", "School ID", "Field", "Raw value", "Normalized value", "Source URL",
                                  "Source title", "Source type", "Publication date", "Retrieved at", "Excerpt",
                                  "Verification status", "Confidence", "Extraction method", "Superseded"],
                {4: 30, 5: 30, 6: 50, 11: 70})
    for r in rk:
        for e in r["evidence"]:
            _append(ws, [e["evidence_id"], e["school_id"], e["field_name"], e["raw_value"][:1000],
                       _fmt_values(e["normalized_value"]), None, e["source_title"], e["source_type"], e.get("publication_date"),
                       e["retrieved_at"], e["evidence_excerpt"][:1500], e["verification_status"], e["confidence_score"],
                       e["extraction_method"], e.get("superseded", False)])
            _link(ws, ws.max_row, 6, e["source_url"])

    # 12. Source Registry ------------------------------------------------------------------------------------
    ws = _sheet(wb, "Source Registry", ["Source", "URL", "Tier", "Responsibility", "Access this run", "Reason", "Checked at"],
                {1: 30, 2: 45, 4: 45, 6: 60})
    for s in sources:
        _append(ws, [s["name"], None, s["tier"], s.get("responsibility"), s.get("access"), s.get("reason"), s.get("checked_at")])
        _link(ws, ws.max_row, 2, s.get("url"))
    _wrap(ws)

    out = cfg.paths.output_dir / cfg.research["report"]["filename"]
    out.parent.mkdir(parents=True, exist_ok=True)
    tmp = out.with_suffix(".tmp.xlsx")
    wb.save(tmp)
    tmp.replace(out)
    return str(out)


def _fmt_values(v) -> str | None:
    if v is None:
        return None
    if isinstance(v, (str, int, float)):
        return str(v)
    import json

    return json.dumps(v, ensure_ascii=False)[:2000]


__all__ = ["build_report", "grade_label", "VERIFIED_FIELD_STATUSES"]
